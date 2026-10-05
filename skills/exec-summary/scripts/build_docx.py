#!/usr/bin/env python3
"""Render an executive-summary markdown file as a .docx in the Innovation Team template's look.

Calibri throughout, teal 15697B headings and rules, pale-teal table header rows,
sand callout boxes, red for any [bracketed placeholder] still left in the text.
The supplied Word template carries no reusable named styles (it is all direct
formatting), so the formatting is applied directly here for the same reason.

Usage:
  python3 build_docx.py summary.md out.docx [--images DIR]

Needs python-docx. Without it installed:
  uv run --with python-docx python3 build_docx.py summary.md out.docx

Markdown the builder understands (see references/skeleton.md):
  cover block before the first \\newpage line:
    # CODENAME / ## descriptor / ### scope line / **Executive Summary**
    then blank-line separated blocks: author blocks (name, role line, country),
    a month line, and a "Classification: ..." line
  body: ## and ### headings, | tables |, - bullets, 1. numbered items,
    > callout boxes, ![caption](file.png), \\newpage, and a closing
    "Document Version: x" line that also feeds the footer.
"""
import argparse
import os
import re
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.image.image import Image as DocxImage
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is not installed. Run: uv run --with python-docx python3 "
             + " ".join(sys.argv) + "   (or pip install python-docx)")

TEAL = RGBColor(0x15, 0x69, 0x7B)
INK = RGBColor(0x1F, 0x2A, 0x33)
MUTE = RGBColor(0x5D, 0x66, 0x73)
RED = RGBColor(0xC0, 0x1F, 0x1F)
HDRFILL, ROWFILL, SANDFILL = "EDF4F6", "F4F7F8", "F7F5EF"
FONT, TEXTW = "Calibri", Inches(6.5)
NEWPAGE = "\\newpage"


def shade(cell, hexfill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexfill)
    cell._tc.get_or_add_tcPr().append(el)


def box_borders(table, fill):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "15697B" if edge == "left" else fill)
        borders.append(e)
    table._tbl.tblPr.append(borders)


def runfmt(r, size=10.5, bold=False, color=INK, italic=False):
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), FONT)
    r._element.get_or_add_rPr().append(rf)


# Flat inline markers only: **bold**, *italic*, `code`, [placeholder]. Markdown
# links are not used in these documents, so a bare [..] is always a placeholder.
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+\])")


def add_rich(par, text, size=10.5, color=INK, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runfmt(par.add_run(part[2:-2]), size, True, color)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runfmt(par.add_run(part[1:-1]), size, bold, color, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            runfmt(r, size - 0.5, bold, MUTE)
            r.font.name = "Consolas"
        elif part.startswith("[") and part.endswith("]"):
            runfmt(par.add_run(part), size, bold, RED)
        else:
            runfmt(par.add_run(part), size, bold, color)


def rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "24"), ("w:color", "15697B"), ("w:space", "1")):
        bt.set(qn(k), v)
    b.append(bt)
    p._p.get_or_add_pPr().append(b)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_field(par, instr):
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        r = par.add_run()
        runfmt(r, 8.5, False, MUTE)
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)


def parse_cover(cover_src):
    """Blank-line separated blocks. Heading lines carry codename, descriptor, scope."""
    blocks, cur = [], []
    for line in cover_src.split("\n"):
        if line.strip():
            cur.append(line.strip())
        elif cur:
            blocks.append(cur)
            cur = []
    if cur:
        blocks.append(cur)
    cover = {"codename": "", "descriptor": "", "scope": "", "authors": [], "month": "",
             "classification": ""}
    for blk in blocks:
        for line in list(blk):
            if line.startswith("### "):
                cover["scope"] = line[4:].strip()
            elif line.startswith("## "):
                cover["descriptor"] = line[3:].strip()
            elif line.startswith("# "):
                cover["codename"] = line[2:].strip()
            elif line.lower().startswith("classification"):
                cover["classification"] = line
            elif line.strip("*").strip().lower() == "executive summary":
                pass
            else:
                continue
            blk.remove(line)
        if not blk:
            continue
        if len(blk) == 1 and not cover["month"]:
            cover["month"] = blk[0]
        elif len(blk) == 2 and blk[1].lower().startswith("classification"):
            cover["month"], cover["classification"] = blk
        else:
            cover["authors"].append(blk)
    # month and classification are often written as one two-line block
    for a in list(cover["authors"]):
        if len(a) == 2 and a[1].lower().startswith("classification"):
            cover["month"], cover["classification"] = a
            cover["authors"].remove(a)
    return cover


def build(md_path, out_path, img_dir):
    src = open(md_path, encoding="utf-8").read()
    if NEWPAGE not in src:
        sys.exit("no \\newpage line after the cover block. See references/skeleton.md")
    body_start = src.index(NEWPAGE)
    cover = parse_cover(src[:body_start])
    version = ""
    for line in src.splitlines():
        if line.strip().lower().startswith("document version:"):
            version = line.split(":", 1)[1].strip()

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.9)
        s.left_margin = s.right_margin = Inches(1.0)
        s.different_first_page_header_footer = True
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label = " · ".join(x for x in (
            "Project " + cover["codename"].replace("PROJECT ", "").title() if cover["codename"] else "",
            "Executive Summary", ("Version " + version) if version else "", cover["month"]) if x)
        add_rich(fp, label + " · Page ", 8.5, MUTE)
        add_field(fp, "PAGE")
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.08
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-GB")
    st.element.get_or_add_rPr().append(lang)

    def cov(txt, size, bold=True, color=TEAL, before=0, after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        add_rich(p, txt, size, color, bold)
        return p

    # Cover page is built, not flowed: its whole first impression is vertical rhythm.
    cov(cover["codename"] or "[PROJECT CODENAME]", 40, True, TEAL, 96, 4)
    cov(cover["descriptor"] or "[One-line product descriptor]", 22, True, INK, 0, 10)
    rule(doc)
    cov(cover["scope"] or "[Scope line]", 12.5, False, MUTE, 0, 30)
    cov("Executive Summary", 13, True, TEAL, 0, 26)
    for a in cover["authors"] or [["[Author name]", "[Role, programme, business unit]", "[Country]"]]:
        cov(a[0], 11.5, True, INK, 0, 1)
        for extra in a[1:-1]:
            cov(extra, 10.5, False, MUTE, 0, 1)
        if len(a) > 1:
            cov(a[-1], 10.5, False, MUTE, 0, 14)
    cov(cover["month"] or "[Month Year]", 10.5, False, MUTE, 26, 2)
    cov(cover["classification"] or "Classification: [Internal / Commercial in Confidence]",
        10.5, True, TEAL, 0, 0)
    page_break(doc)

    lines = src[body_start + len(NEWPAGE):].split("\n")
    pending = []

    def flush_table():
        if not pending:
            return
        rows = [r for r in pending if not set(r.strip()) <= set("|:- ")]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        ncol = max(len(c) for c in cells)
        t = doc.add_table(rows=0, cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        for ri, row in enumerate(cells):
            row += [""] * (ncol - len(row))
            cs = t.add_row().cells
            for ci, val in enumerate(row):
                cell = cs[ci]
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                add_rich(p, val, 9, TEAL if ri == 0 else INK, ri == 0)
                shade(cell, HDRFILL if ri == 0 else (ROWFILL if ri % 2 else "FFFFFF"))
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        del pending[:]

    for ln in lines:
        if ln.startswith("|"):
            pending.append(ln)
            continue
        flush_table()
        s = ln.strip()
        if s == NEWPAGE:
            page_break(doc)
        elif s.startswith("!["):
            cap = s[2:s.index("](")] if "](" in s else ""
            fn = s[s.index("](") + 2:s.rindex(")")] if "](" in s else ""
            path = os.path.join(img_dir, fn) if img_dir else fn
            if not os.path.isfile(path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_rich(p, "[Insert image: %s]" % (cap or fn), 10.5)
                continue
            img = DocxImage.from_file(path)
            ratio = img.px_width / float(img.px_height or 1)
            # wide screenshots get full width, tall ones are held back so they do not eat the page
            wid = TEXTW if ratio >= 1.7 else Inches(4.6) if ratio >= 1.1 else Inches(3.4)
            doc.add_picture(path, width=wid)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c = doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.space_after = Pt(12)
            runfmt(c.add_run(cap), 8.5, False, MUTE, italic=True)
        elif s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#"))
            txt = s.lstrip("# ").strip()
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(14 if lvl <= 2 else 10)
            pf.space_after = Pt(4)
            pf.keep_with_next = True
            sizes = {1: 24, 2: 16, 3: 12.5, 4: 11}
            add_rich(p, txt, sizes.get(lvl, 11), TEAL if lvl <= 3 else INK, True)
            if lvl == 2:
                rule(doc)
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            add_rich(p, s[2:])
        elif s[:1].isdigit() and ". " in s[:4]:
            # numbers written as text: Word's List Number style shares one sequence
            # across the document, so a second list would start at 4
            num, rest = s.split(". ", 1)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            add_rich(p, num + ".\t" + rest)
        elif s == "---":
            rule(doc)
        elif s.startswith("> "):
            t = doc.add_table(rows=1, cols=1)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = t.rows[0].cells[0]
            cell.text = ""
            shade(cell, SANDFILL)
            box_borders(t, SANDFILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            add_rich(p, s[2:], 10)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif s:
            add_rich(doc.add_paragraph(), s)
    flush_table()
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("markdown")
    ap.add_argument("out")
    ap.add_argument("--images", default=None, help="directory holding the ![..](file) images")
    a = ap.parse_args()
    img_dir = a.images or os.path.dirname(os.path.abspath(a.markdown))
    print(build(a.markdown, a.out, img_dir))
